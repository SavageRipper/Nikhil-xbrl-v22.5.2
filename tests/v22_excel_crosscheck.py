
#!/usr/bin/env python3
import sys
from pathlib import Path
from openpyxl import load_workbook
if len(sys.argv)!=2:
    raise SystemExit("Usage: python3 tests/v22_excel_crosscheck.py /path/to/CompuXBRL-workbook.xlsx")
xlsx=Path(sys.argv[1]); root=Path(__file__).resolve().parents[1]
cat=(root/'TAXONOMY_TABLE_CATALOG.csv').read_text(encoding='utf-8')
wb=load_workbook(xlsx,data_only=False,read_only=True)
table_sheets=[s for s in wb.sheetnames if wb[s]['A2'].value=='table']
unique={str(wb[s]['K2'].value) for s in table_sheets if wb[s]['K2'].value}
assert len(table_sheets)==134,len(table_sheets)
assert len(unique)==87,len(unique)
for sheet,q in {
 'LoansAndAdvances':'LoansAndAdvancesTable',
 'BreakupOfProvisions':'DisclosureOfBreakupOfProvisionsTable',
 'RawMaterialsConsumed':'DetailsOfRawMaterialsConsumedTable',
 'GoodsPurchased':'DetailsOfGoodsPurchasedTable',
 'ManufacturedAndTradedGoods':'DetailsOfManufacturedAndTradedGoodsTable',
 'WorkInProgress':'DetailsOfWorkInProgressTable',
}.items():
 assert sheet in wb.sheetnames,sheet
 assert str(wb[sheet]['K2'].value)==q,(sheet,wb[sheet]['K2'].value)
# The workbook's table sheets expose the same [Table] -> line-item structure used by V22.
for sheet,headers in {
 'LoansAndAdvances':('DisclosureOfLoansAndAdvancesLineItems','MemberElement1','MemberElement2'),
 'RawMaterialsConsumed':('DetailsOfRawMaterialsConsumedLineItems','DescriptionOfRawMaterialsCategory','RawMaterialsConsumed'),
 'GoodsPurchased':('DetailsOfGoodsPurchasedLineItems','DescriptionOfGoodsPurchased','GoodsPurchased'),
 'ManufacturedAndTradedGoods':('DetailsOfManufacturedAndTradedGoodsLineItems','CustomElement1','DescriptionOfFinishedGoods'),
 'WorkInProgress':('DetailsOfWorkInProgressLineItems','DescriptionOfWorkInProgress','WorkInProgress'),
}.items():
 ws=wb[sheet]
 assert ws['K3'].value==headers[0],(sheet,ws['K3'].value)
 assert ws['M3'].value==headers[1],(sheet,ws['M3'].value)
 assert ws['N3'].value==headers[2],(sheet,ws['N3'].value)
print(f"PASS: V22 CompuXBRL workbook cross-check ({len(table_sheets)} table sheets, {len(unique)} unique table identifiers)")
